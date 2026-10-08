#!/usr/bin/env python3
"""Convert the maintained report Markdown and actual build images to a Word draft.

Usage: /workspace/.runner-tools/venv/bin/python scripts/make-report.py
Requires python-docx==1.1.2. No course attachments are needed to regenerate.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "DANIYAR_MIDTERM.md"
DEST = ROOT / "docs" / "Daniyar_Midterm_Draft.docx"


def runs(paragraph, text):
    """Keep simple emphasis without exposing Markdown markers in the document."""
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string("166D67")
    header = section.header.paragraphs[0]
    header.text = "AITU  /  COMPUTER GRAPHICS FUNDAMENTALS  /  TOPIC 4"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Daniyar — midterm draft  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Daniyar — Instanced Obstacles Midterm Draft"
    doc.core_properties.subject = "Topic 4: Endless runner with instanced obstacles and a height-map ground"
    doc.core_properties.author = "Daniyar and Muhamedzhan (draft prepared with Codex assistance)"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(lines[i])
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                i += 1
        elif line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            if not match:
                raise ValueError(f"Invalid image line: {line}")
            caption, relative = match.groups()
            path = SOURCE.parent / relative
            # Cap diagram height so it fits on one A4 page with its caption.
            from PIL import Image
            with Image.open(path) as picture:
                width = min(6.75, 5.6 * picture.width / picture.height)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Inches(width))
            doc.add_paragraph(caption, style="Caption")
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", v) for v in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for j, value in enumerate(rows[0]):
                runs(table.rows[0].cells[j].paragraphs[0], value)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                cells = table.add_row().cells
                for j, value in enumerate(row):
                    runs(cells[j].paragraphs[0], value)
            for row in table.rows:
                no_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
            doc.add_paragraph()
            continue
        else:
            runs(doc.add_paragraph(), line)
        i += 1
    doc.save(DEST)
    # Read back the OOXML package and verify its essential content and evidence.
    check = Document(DEST)
    assert len(check.inline_shapes) == 2, "Both real evidence images must be embedded"
    assert len(check.tables) == 4, "Tools, timings, checklist and contribution tables expected"
    assert any("References" in p.text for p in check.paragraphs)
    print(f"Written {DEST}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, 2 images")


if __name__ == "__main__":
    main()
